import streamlit as st
from openai import OpenAI
import os
import json
import html
import urllib.parse
import math
import io
import re
import PyPDF2
import docx
import pandas as pd
import plotly.graph_objects as go

# ================================================================ setup

def get_api_key():
    try:
        return st.secrets["DEEPSEEK_API_KEY"]
    except Exception:
        return os.environ.get("DEEPSEEK_API_KEY")

API_KEY = get_api_key()
MODEL = "deepseek-chat"

def get_client():
    return OpenAI(api_key=API_KEY, base_url="https://api.deepseek.com")

def extract_text_from_file(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        reader = PyPDF2.PdfReader(uploaded_file)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        doc = docx.Document(uploaded_file)
        return "\n".join(p.text for p in doc.paragraphs)
    return ""

# ================================================================ analysis prompt

SYSTEM_PROMPT = """You are a skeptical senior product manager with 15+ years of experience taking early-stage
products from idea to launch. Your job is to find the assumptions the idea depends on, rank the riskiest ones,
and lay out a concrete, cheap validation plan the founder can start this week.

Rules:
- Be critical and specific. No generic advice. Use concrete numbers, names, channels, and sample sizes.
- Label each claim with a "certainty" field, using only two values:
    "Fact"     = the user stated it, or it is widely documented public knowledge
    "Estimate" = your own reasoning or guess; the founder should verify it
- You have no live market data. Any market size, growth, or price figure must be "Estimate"
  unless the user provided it. For market size, show a bottom-up calculation (users x price x frequency).
- Pick exactly 3 riskiest assumptions, most dangerous first. Classify each as Desirability (do people want it),
  Viability (will it make money), Feasibility (can it be built), or Usability (can people use it).
- Validation plan: 3-4 steps ordered from cheapest/fastest to most expensive. Test desirability before
  viability, and both before building anything expensive. Each step names the method, sample size,
  exactly where to recruit participants, the deliverable, and a decision rule (continue / change / stop).
- Competitors: 4-6 real companies or products when possible (direct, indirect, and substitutes). Give typical
  price, key features, and target user. Choose two axes for a positioning map that matter most to buyers,
  and place each competitor and the user's product on 1-10 scales.
- Business model canvas: 2-4 specific bullet points per block.
- Unit economics are rough estimates.
- The score (0-100) is a secondary reference only. Above 70 should be rare.
{mode_rules}

Return ONLY valid json in exactly this structure:
{{
  "one_line_summary": "what the idea is, in plain words",
  "assumed_target_user": "who this is for, as specifically as possible",
  "stage": "Idea | Prototype | MVP | Launched",
  "confidence": "Low | Medium | High",
  "confidence_reason": "two sentences: what is known and what is not",
  "score": 0,
  "riskiest_assumptions": [
    {{
      "assumption": "the belief the idea depends on",
      "risk_type": "Desirability | Viability | Feasibility | Usability",
      "why_risky": "what happens if it is wrong",
      "test_method": "a concrete test",
      "test_time": "e.g. 3 days",
      "sample": "e.g. 8 interviews",
      "pass_signal": "the result that means the assumption holds"
    }}
  ],
  "validation_plan": [
    {{
      "step": "short name",
      "timing": "e.g. Week 1",
      "goal": "which assumption this tests and why now",
      "method": "Interviews | Survey | Landing page | Prototype test | Concierge test | Pilot | Desk research",
      "details": "what exactly to do",
      "sample_size": "e.g. 8-10 people",
      "where_to_recruit": "specific places, e.g. r/SkincareAddiction, campus clubs, LinkedIn groups",
      "deliverable": "what you will have at the end",
      "decision_rule": "continue if..., change if..., stop if..."
    }}
  ],
  "what_would_change_my_mind": {{
    "raise_confidence": ["specific evidence"],
    "lower_confidence": ["specific evidence"]
  }},
  "market": {{
    "tam": {{"value": "e.g. $2B", "calculation": "bottom-up formula", "certainty": "Fact | Estimate", "how_to_verify": "..."}},
    "sam": {{"value": "...", "calculation": "...", "certainty": "Fact | Estimate", "how_to_verify": "..."}},
    "som": {{"value": "...", "calculation": "...", "certainty": "Fact | Estimate", "how_to_verify": "..."}},
    "trends": [{{"claim": "...", "certainty": "Fact | Estimate"}}]
  }},
  "positioning_axes": {{"x_label": "e.g. Price (low to high)", "y_label": "e.g. Routine guidance (none to full)"}},
  "your_position": {{"x": 5, "y": 5}},
  "competitors": [
    {{
      "name": "...", "type": "Direct | Indirect | Substitute", "price": "typical price",
      "key_features": "...", "target_user": "...", "strength": "...", "weakness": "...",
      "x": 5, "y": 5, "certainty": "Fact | Estimate"
    }}
  ],
  "user_pain_points": {{
    "jtbd": [{{"claim": "...", "certainty": "Fact | Estimate"}}],
    "frustrations": [{{"claim": "...", "certainty": "Fact | Estimate"}}]
  }},
  "business_model": {{
    "key_partners": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "key_activities": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "key_resources": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "value_propositions": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "customer_relationships": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "channels": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "customer_segments": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "cost_structure": [{{"point": "...", "certainty": "Fact | Estimate"}}],
    "revenue_streams": [{{"point": "...", "certainty": "Fact | Estimate"}}]
  }},
  "unit_economics": {{
    "price": "...", "unit_cost": "...", "gross_margin": "...",
    "customer_acquisition_cost": "...", "payback": "...", "note": "the biggest unknown here"
  }}
}}"""

B2C_RULES = "- Treat this as a consumer product: the user and the buyer are usually the same person."
B2B_RULES = """- Treat this as a B2B product. Separate the BUYER (who pays and signs) from the USER (who uses it daily).
- At least one riskiest assumption must be about the buyer: budget owner, procurement, switching cost, or sales cycle.
- Consider compliance, security review, and integration with existing systems."""

def parse_json(text):
    clean = text.strip()
    if clean.startswith("```"):
        clean = clean.split("```")[1]
        if clean.startswith("json"):
            clean = clean[4:]
    return json.loads(clean.strip())

def analyze_product(product_input, mode, previous=None):
    system_prompt = SYSTEM_PROMPT.format(mode_rules=B2B_RULES if mode == "B2B" else B2C_RULES)
    user_msg = "Analyze this product concept and return json only:\n\n" + product_input
    if previous:
        user_msg += ("\n\nThis is a revised version. The previous riskiest assumptions were:\n- "
                     + "\n- ".join(a.get("assumption", "") for a in previous.get("riskiest_assumptions", []))
                     + "\nRe-rank them for the revised idea.")
    response = get_client().chat.completions.create(
        model=MODEL, max_tokens=7000, temperature=0.2,
        response_format={"type": "json_object"},
        messages=[{"role": "system", "content": system_prompt},
                  {"role": "user", "content": user_msg}])
    text = response.choices[0].message.content
    try:
        return parse_json(text)
    except Exception:
        return {"error": text}

# ================================================================ test kit prompt

KIT_PROMPT = """You are a senior UX researcher and market research analyst. Write ready-to-use research materials
for a founder who will run this validation step tomorrow. Be concrete and specific to the product.
Rules:
- Interview questions ask about past behavior and real situations, never "would you use this?" or leading questions.
- Surveys: start with 2-3 screener questions, use clear single-topic questions, include answer options.
  If price is being tested, include the four Van Westendorp price sensitivity questions.
  End with a short "How to analyze the results" section.
- Recruiting posts are short, honest about time needed, and say what participants get.
- Landing page copy: headline, subheadline, 3 benefits, call to action, price line, and what metric to track.
Write in clear English using Markdown headings. Only include the sections requested."""

KIT_OPTIONS = {
    "Screener questions": "Screener questions to find the right participants",
    "Interview guide": "A 30-minute interview guide with intro script, 10-12 questions, and probes",
    "Survey": "A survey of 12-15 questions",
    "Recruiting post": "A recruiting post for the places named in the plan",
    "Landing page copy": "Landing page copy for a smoke test",
    "Prototype test script": "A usability test script with 3-5 tasks and success criteria",
}

def default_kit(method):
    m = (method or "").lower()
    if "interview" in m:
        return ["Screener questions", "Interview guide", "Recruiting post"]
    if "survey" in m:
        return ["Survey", "Recruiting post"]
    if "landing" in m:
        return ["Landing page copy", "Recruiting post"]
    if "prototype" in m or "usability" in m:
        return ["Screener questions", "Prototype test script", "Recruiting post"]
    return ["Screener questions", "Interview guide"]

def generate_kit(result, step, parts):
    brief = {
        "product": result.get("one_line_summary"),
        "target_user": result.get("assumed_target_user"),
        "riskiest_assumptions": [a.get("assumption") for a in result.get("riskiest_assumptions", [])],
        "validation_step": step,
        "sections_requested": [KIT_OPTIONS[p] for p in parts],
    }
    response = get_client().chat.completions.create(
        model=MODEL, max_tokens=5000, temperature=0.4,
        messages=[{"role": "system", "content": KIT_PROMPT},
                  {"role": "user", "content": json.dumps(brief, ensure_ascii=False, indent=2)}])
    return response.choices[0].message.content

# ================================================================ PRD / MRD

MRD_TEMPLATE = """# Market Requirements Document: [Product name]
## 1. Executive summary
## 2. Market problem
## 3. Target market and segments
## 4. Market size (TAM / SAM / SOM, with calculation)
## 5. Customer needs and jobs-to-be-done
## 6. Competitive landscape
## 7. Positioning and differentiation
## 8. Pricing and business model
## 9. Go-to-market approach
## 10. Success criteria
## 11. Assumptions and open questions
## 12. Validation plan"""

PRD_TEMPLATE = """# Product Requirements Document: [Product name]
## 1. Overview (problem, solution, why now)
## 2. Goals and non-goals
## 3. Target users and personas
## 4. User stories and key use cases
## 5. Requirements (table: ID, requirement, priority P0/P1/P2, rationale, acceptance criteria)
## 6. MVP scope (in scope / deferred, with reasons)
## 7. User flow
## 8. Success metrics (North Star, supporting metrics, guardrail metrics)
## 9. Dependencies and technical considerations
## 10. Risks and mitigations
## 11. Assumptions and open questions
## 12. Release plan and milestones"""

DOC_PROMPT = """You are a senior product manager writing a {doc_type} for an early-stage product.
Write it from the analysis and inputs provided. Be specific and concrete; no filler.
Rules:
- Follow the template's section structure exactly, in the same order, using Markdown headings.
- Never invent research findings, quotes, or user counts. Where real research is needed but not provided,
  write a clear placeholder in square brackets, e.g. [Add findings from 8 user interviews].
- Mark any statement that depends on an untested assumption with "(to validate)".
- Requirements: prioritize as P0 / P1 / P2 with a one-line rationale and testable acceptance criteria,
  and keep P0 to the smallest set that tests the riskiest assumptions.
- Metrics: one North Star metric, 2-3 supporting metrics, 1-2 guardrail metrics, each with a definition.
- If test results are provided, use them, and move any assumption they confirm out of open questions.
- Use Markdown tables where the template asks for a table.
Template:
{template}"""

def generate_doc(result, idea, doc_type, template, test_results):
    payload = {
        "original_description": idea,
        "analysis": result,
        "test_results_so_far": test_results or "None yet",
    }
    response = get_client().chat.completions.create(
        model=MODEL, max_tokens=7000, temperature=0.3,
        messages=[{"role": "system", "content": DOC_PROMPT.format(doc_type=doc_type, template=template)},
                  {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}])
    return response.choices[0].message.content

def add_runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part:
            par.add_run(part)

def markdown_to_docx(md):
    d = docx.Document()
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in block
                    if not re.match(r"^\|?\s*:?-{2,}", r.replace("|", "", 1))]
            if rows:
                t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci, c in enumerate(r):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        add_runs(cell.paragraphs[0], c)
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 3) if len(m.group(1)) > 1 else 0)
        elif re.match(r"^\s*[-*]\s+", line):
            add_runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.strip() and not re.match(r"^-{3,}$", line.strip()):
            add_runs(d.add_paragraph(), line.strip())
        i += 1
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

def read_template(file):
    name = file.name.lower()
    if name.endswith(".docx"):
        d = docx.Document(file)
        out = []
        for p in d.paragraphs:
            if not p.text.strip():
                continue
            if p.style.name.lower().startswith("heading") or p.style.name.lower() == "title":
                out.append("## " + p.text.strip())
            else:
                out.append(p.text.strip())
        return "\n".join(out)
    if name.endswith(".pdf"):
        return extract_text_from_file(file)
    return file.read().decode("utf-8", errors="ignore")

# ================================================================ display helpers

def norm_certainty(v):
    t = str(v or "").lower()
    if any(k in t for k in ("fact", "back", "evid")):
        return "Fact"
    if any(k in t for k in ("estim", "guess", "infer", "untest", "assum")):
        return "Estimate"
    return None

def esc(x):
    return html.escape(str(x if x is not None else ""))

def tag(v):
    # Only estimates are marked, so the page stays calm and the mark means "check this".
    if norm_certainty(v) != "Estimate":
        return ""
    return ('<span title="The AI\'s best guess. Check it before you rely on it." '
            'style="border:1px solid #c98a00;color:#c98a00;border-radius:10px;padding:0 7px;'
            'font-size:11px;font-weight:600;margin-left:6px;white-space:nowrap;cursor:help;">Estimate</span>')

def label_text(v):
    return {"Fact": "Fact", "Estimate": "Estimate — verify"}.get(norm_certainty(v), "—")

def legend():
    st.caption("Points marked " + "Estimate" + " are the AI's best guess. Check them before relying on them. "
               "Unmarked points come from your description or well-known public information.")

def claim_lines(items, field="claim"):
    out = []
    for item in items or []:
        if isinstance(item, dict):
            out.append(f'<div style="margin:6px 0;">• {esc(item.get(field, ""))}{tag(item.get("certainty"))}</div>')
        else:
            out.append(f'<div style="margin:6px 0;">• {esc(item)}</div>')
    return "".join(out)

def ab_sample_size(p1, p2, z_alpha=1.96, z_beta=0.8416):
    """Visitors needed per variant to detect p1 vs p2 (two-sided 5% significance, 80% power)."""
    if p1 <= 0 or p2 <= 0 or p1 >= 1 or p2 >= 1 or p1 == p2:
        return None
    pbar = (p1 + p2) / 2
    num = (z_alpha * math.sqrt(2 * pbar * (1 - pbar)) + z_beta * math.sqrt(p1 * (1 - p1) + p2 * (1 - p2))) ** 2
    return math.ceil(num / (p1 - p2) ** 2)

def search_url(name):
    return "https://www.google.com/search?q=" + urllib.parse.quote_plus(str(name))

def to_num(v, default=5.0):
    try:
        return max(0.0, min(10.0, float(v)))
    except Exception:
        return default

# ================================================================ page

st.set_page_config(page_title="PM Product Analyzer", page_icon="📊", layout="wide")
st.title("📊 PM Product Analyzer")
st.caption("Enter a product concept. Get its riskiest assumptions, a step-by-step validation plan, "
           "ready-to-use research materials, and a market, competitor, and business model overview.")

for k, v in {"history": [], "kit": None, "doc": None}.items():
    st.session_state.setdefault(k, v)

if not API_KEY:
    st.error("DEEPSEEK_API_KEY is not set. In Streamlit Cloud, open Manage app → Settings → Secrets, "
             "add DEEPSEEK_API_KEY = \"sk-...\", save, then reboot the app.")

with st.expander("📋 Suggested input format"):
    st.markdown("""
Fill in as many fields as possible for a more accurate analysis:

```
Product name: [name]
Target users: [who are they, age range, role]
Core problem: [what pain point are you solving]
Solution: [how your product solves it]
Market: [which market/industry]
Monetization: [how you plan to make money]
Stage: [idea / MVP / prototype]
```
""")

st.subheader("Enter Product Information")
mode = st.radio("Business type", ["B2C", "B2B"], horizontal=True,
                help="B2B separates the buyer (who pays) from the user (who uses it).")

col1, col2 = st.columns(2)
with col1:
    text_input = st.text_area(
        "Text Description",
        placeholder="Product name: ...\nTarget users: ...\nCore problem: ...\nSolution: ...\nMarket: ...\nMonetization: ...",
        height=220, key="text_input")
with col2:
    uploaded_file = st.file_uploader("Or upload a PDF / DOCX", type=["pdf", "docx"],
                                     key=f"uploader_{st.session_state.get('uploader_key', 0)}")
    file_text = ""
    if uploaded_file:
        file_text = extract_text_from_file(uploaded_file)
        st.success(f"File read successfully — {len(file_text)} characters extracted.")

final_input = (text_input or "") + ("\n\n" + file_text if file_text else "")

def clear_inputs():
    st.session_state["text_input"] = ""
    st.session_state["uploader_key"] = st.session_state.get("uploader_key", 0) + 1
    st.session_state["history"] = []
    st.session_state["kit"] = None
    st.session_state["doc"] = None

disabled = not final_input.strip() or not API_KEY
update = False
if st.session_state["history"]:
    st.caption("Changed the same idea? Choose Update to compare with the last result. "
               "Working on a different product? Choose New idea.")
    col_btn1, col_btn2, col_btn3, col_rest = st.columns([1.6, 1.6, 1, 4])
    with col_btn1:
        update = st.button("Update this idea", type="primary", disabled=disabled)
    with col_btn2:
        new_idea = st.button("Analyze as a new idea", disabled=disabled)
    with col_btn3:
        st.button("Clear", type="secondary", on_click=clear_inputs)
    if new_idea:
        st.session_state["history"] = []
    analyze = update or new_idea
else:
    col_btn1, col_btn2, col_rest = st.columns([1.5, 1, 6])
    with col_btn1:
        analyze = st.button("Analyze Product", type="primary", disabled=disabled)
    with col_btn2:
        st.button("Clear", type="secondary", on_click=clear_inputs)

if analyze:
    previous = st.session_state["history"][-1]["result"] if st.session_state["history"] else None
    with st.spinner("Analyzing your product concept... this may take 30-60 seconds."):
        try:
            result = analyze_product(final_input, mode, previous)
        except Exception as e:
            result = {"error": f"Request failed: {e}"}
    if "error" in result:
        st.error("Analysis failed. Click Analyze Product again. Details:")
        st.code(result["error"])
    else:
        st.session_state["history"].append({"idea": final_input, "mode": mode, "result": result})
        st.session_state["kit"] = None
        st.session_state["doc"] = None

if not st.session_state["history"]:
    st.stop()

# ================================================================ results

result = st.session_state["history"][-1]["result"]
prev = st.session_state["history"][-2]["result"] if len(st.session_state["history"]) > 1 else None

if prev:
    st.success("Analysis updated. The Validation Plan tab compares it with your previous description.")
else:
    st.success("Analysis complete!")
st.divider()

tab1, tab2, tab3, tab4, tab5, tab6, tab7 = st.tabs([
    "🎯 Riskiest Assumptions", "🧪 Validation Plan", "🏆 Competitors",
    "📈 Market Size", "👤 User Pain Points", "💰 Business Model", "📄 MRD / PRD"
])

# ---------------- tab 1: assumptions
with tab1:
    st.markdown(f"### {esc(result.get('one_line_summary', ''))}")
    st.caption(f"Assumed target user: {result.get('assumed_target_user', '')}  |  Stage: {result.get('stage', '—')}")
    st.write("These are the beliefs your idea depends on that nobody has tested yet, most dangerous first. "
             "Test them before investing in building.")
    rows = []
    for i, a in enumerate(result.get("riskiest_assumptions", []), 1):
        rows.append({"#": i, "Assumption": a.get("assumption", ""), "Type": a.get("risk_type", ""),
                     "Test": a.get("test_method", ""),
                     "Time": a.get("test_time", ""), "Sample": a.get("sample", ""),
                     "It holds if": a.get("pass_signal", "")})
    if rows:
        st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)
    st.markdown("#### Why each one matters")
    for i, a in enumerate(result.get("riskiest_assumptions", []), 1):
        with st.expander(f"{i}. {a.get('assumption', '')}"):
            st.write(f"**Type:** {a.get('risk_type', '')}")
            st.write(f"**If this is wrong:** {a.get('why_risky', '')}")
    st.caption("Desirability = do people want it. Viability = will it make money. "
               "Feasibility = can it be built. Usability = can people use it.")

# ---------------- tab 2: validation plan + test kit
with tab2:
    conf = str(result.get("confidence", "")).split("|")[0].strip().title()
    try:
        score = int(float(result.get("score", 0)))
    except Exception:
        score = 0
    banner = f"**Confidence: {conf}** — {result.get('confidence_reason', '')}"
    if conf == "High":
        st.success(banner)
    elif conf == "Medium":
        st.info(banner)
    else:
        st.warning(banner)
    st.caption(f"Reference score: {score} / 100. Confidence reflects how much is already known, "
               "not how good the idea is. The plan below is how you raise it.")
    if prev:
        st.info(f"Compared with your previous description: confidence {prev.get('confidence', '?')} → {conf}, "
                f"score {prev.get('score', '?')} → {score}")

    st.markdown("### Your validation plan")
    st.write("Cheapest and fastest tests first. Only move to the next step if the decision rule says continue.")
    plan = result.get("validation_plan", [])
    for i, s in enumerate(plan, 1):
        with st.container(border=True):
            st.markdown(f"**Step {i} · {esc(s.get('timing', ''))}: {esc(s.get('step', ''))}**")
            st.write(f"**Goal:** {s.get('goal', '')}")
            c1, c2, c3 = st.columns(3)
            c1.markdown(f"**Method**  \n{s.get('method', '')}")
            c2.markdown(f"**Sample size**  \n{s.get('sample_size', '')}")
            c3.markdown(f"**You'll end up with**  \n{s.get('deliverable', '')}")
            st.write(f"**What to do:** {s.get('details', '')}")
            st.write(f"**Where to recruit:** {s.get('where_to_recruit', '')}")
            st.write(f"**Decision rule:** {s.get('decision_rule', '')}")

    w = result.get("what_would_change_my_mind", {})
    ca, cb = st.columns(2)
    with ca:
        st.markdown("**Results that would raise confidence:**")
        for x in w.get("raise_confidence", []):
            st.write(f"- {x}")
    with cb:
        st.markdown("**Results that would lower confidence:**")
        for x in w.get("lower_confidence", []):
            st.write(f"- {x}")

    if plan:
        st.divider()
        st.markdown("### Build a test kit")
        st.write("Pick a step and get ready-to-use materials: screener, interview guide, survey, recruiting post, and more.")
        labels = [f"Step {i}: {s.get('step', '')} ({s.get('method', '')})" for i, s in enumerate(plan, 1)]
        idx = st.selectbox("Which step are you running?", range(len(plan)), format_func=lambda i: labels[i])
        parts = st.multiselect("What do you need?", list(KIT_OPTIONS.keys()),
                               default=default_kit(plan[idx].get("method")), key=f"parts_{idx}")
        if st.button("Generate test kit", type="primary", disabled=not parts):
            with st.spinner("Writing your research materials... about 30 seconds."):
                try:
                    st.session_state["kit"] = {"step": labels[idx], "text": generate_kit(result, plan[idx], parts)}
                except Exception as e:
                    st.error(f"Request failed: {e}")
        kit = st.session_state.get("kit")
        if kit:
            with st.container(border=True):
                st.caption(kit["step"])
                st.markdown(kit["text"])
            st.download_button("Download test kit", kit["text"], file_name="test_kit.md", mime="text/markdown")

    st.divider()
    with st.expander("📐 A/B test calculator: do I have enough visitors?"):
        st.write("Before splitting landing page traffic into version A and B, check whether your traffic "
                 "is enough to see a real difference.")
        a1, a2, a3 = st.columns(3)
        base = a1.number_input("Version A conversion (%)", 0.1, 99.0, 5.0, 0.5)
        target = a2.number_input("Version B conversion you hope for (%)", 0.1, 99.0, 8.0, 0.5)
        traffic = a3.number_input("Visitors you can get in total", 10, 1_000_000, 500, 50)
        n = ab_sample_size(base / 100, target / 100)
        if n:
            st.markdown(f"You need about **{n:,} visitors per version** ({2*n:,} in total) "
                        "to detect this difference with 95% confidence and 80% power.")
            if traffic >= 2 * n:
                st.success("You have enough traffic to run this A/B test.")
            else:
                st.warning(f"You only have {traffic:,} visitors. Test one version first, or only A/B test "
                           "a bigger change (for example, two very different prices).")
        else:
            st.info("Enter two different conversion rates.")

    if prev:
        with st.expander("Compare riskiest assumptions with your previous description"):
            pa, pb = st.columns(2)
            with pa:
                st.markdown("**Previous**")
                for x in prev.get("riskiest_assumptions", []):
                    st.write(f"- {x.get('assumption', '')}")
            with pb:
                st.markdown("**Now**")
                for x in result.get("riskiest_assumptions", []):
                    st.write(f"- {x.get('assumption', '')}")

# ---------------- tab 3: competitors
with tab3:
    comps = result.get("competitors", [])
    if comps:
        st.markdown("### Competitor comparison")
        st.caption("Prices and features are typical estimates. Use the search link to check current details.")
        df = pd.DataFrame([{
            "Name": c.get("name", ""), "Type": c.get("type", ""), "Price": c.get("price", ""),
            "Key features": c.get("key_features", ""), "Target user": c.get("target_user", ""),
            "Strength": c.get("strength", ""), "Weakness": c.get("weakness", ""),
            "Source": label_text(c.get("certainty")), "Look up": search_url(c.get("name", "")),
        } for c in comps])
        st.dataframe(df, hide_index=True, use_container_width=True,
                     column_config={"Look up": st.column_config.LinkColumn("Look up", display_text="Search")})

        axes = result.get("positioning_axes", {})
        you = result.get("your_position", {})
        fig = go.Figure()
        fig.add_trace(go.Scatter(
            x=[to_num(c.get("x")) for c in comps], y=[to_num(c.get("y")) for c in comps],
            mode="markers+text", text=[c.get("name", "") for c in comps], textposition="top center",
            marker=dict(size=13, color="#3b8bd4"), name="Competitors"))
        fig.add_trace(go.Scatter(
            x=[to_num(you.get("x"))], y=[to_num(you.get("y"))], mode="markers+text",
            text=["Your product"], textposition="top center",
            marker=dict(size=18, color="#f0a500", symbol="star"), name="Your product"))
        fig.add_hline(y=5, line_color="#ccc"); fig.add_vline(x=5, line_color="#ccc")
        fig.update_layout(height=520, margin=dict(l=20, r=20, t=30, b=20), showlegend=False,
                          xaxis=dict(title=axes.get("x_label", ""), range=[0, 10.5]),
                          yaxis=dict(title=axes.get("y_label", ""), range=[0, 10.5]))
        st.markdown("### Positioning map")
        st.caption("Positions are estimated by the model. Use it to spot open space, then confirm with users.")
        st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("No competitor data available.")

# ---------------- tab 4: market
with tab4:
    m = result.get("market", {})
    st.markdown("### Market Size Overview")
    st.caption("No live market data is used. These are bottom-up estimates. Check the calculation and verify the inputs.")
    cols = st.columns(3)
    specs = [("TAM", "Total Addressable Market", "tam", "#1a6eb5"),
             ("SAM", "Serviceable Addressable Market", "sam", "#1d9e75"),
             ("SOM", "Serviceable Obtainable Market", "som", "#3b8bd4")]
    for col, (abbr, full, key, bg) in zip(cols, specs):
        item = m.get(key, {})
        if not isinstance(item, dict):
            item = {"value": item}
        with col:
            st.markdown(f"""<div style="background:{bg};border-radius:16px;padding:28px 16px;text-align:center;color:white;">
<div style="font-size:13px;font-weight:600;letter-spacing:1px;opacity:0.85;">{abbr}</div>
<div style="font-size:11px;opacity:0.7;margin-bottom:8px;">{full}</div>
<div style="font-size:28px;font-weight:700;">{esc(item.get("value", "N/A"))}</div></div>""", unsafe_allow_html=True)
            if item.get("calculation"):
                st.markdown(f"**How it's calculated:** {item.get('calculation')}")
            if item.get("how_to_verify"):
                st.caption(f"Verify: {item.get('how_to_verify')}")
    st.markdown("### Key Market Trends")
    legend()
    st.markdown(claim_lines(m.get("trends", [])), unsafe_allow_html=True)

# ---------------- tab 5: user pain
with tab5:
    p = result.get("user_pain_points", {})
    legend()
    st.markdown("**Jobs-to-be-done:**")
    st.markdown(claim_lines(p.get("jtbd", [])), unsafe_allow_html=True)
    st.markdown("**Frustrations with existing solutions:**")
    st.markdown(claim_lines(p.get("frustrations", [])), unsafe_allow_html=True)

# ---------------- tab 6: business model
with tab6:
    b = result.get("business_model", {})
    st.markdown("### Business Model Canvas")
    legend()
    cell_style = "border:0.5px solid #ddd;border-radius:8px;padding:12px;background:#fff;height:100%;box-sizing:border-box;"
    label_style = "font-size:11px;font-weight:600;color:#666;margin-bottom:6px;"
    val_style = "font-size:12px;color:#222;line-height:1.5;overflow-wrap:break-word;word-break:break-word;"
    bot_style = "border:0.5px solid #ddd;border-radius:8px;padding:12px;background:#f8f8f8;box-sizing:border-box;"

    def cell(icon, label, key, style=cell_style):
        v = b.get(key, [])
        if isinstance(v, str):
            v = [{"point": v}]
        if isinstance(v, dict):
            v = [{"point": v.get("text", ""), "certainty": v.get("certainty") or v.get("basis")}]
        body = claim_lines(v, field="point") or "—"
        return (f'<div style="{style}"><div style="{label_style}">{icon} {label}</div>'
                f'<div style="{val_style}">{body}</div></div>')

    st.markdown(f"""
<table style="width:100%;border-collapse:separate;border-spacing:6px;table-layout:fixed;">
<tr>
  <td rowspan="2" style="width:20%;vertical-align:top;">{cell("🔗","Key Partners","key_partners")}</td>
  <td style="width:20%;vertical-align:top;">{cell("✅","Key Activities","key_activities")}</td>
  <td rowspan="2" style="width:20%;vertical-align:top;">{cell("🎁","Value Propositions","value_propositions")}</td>
  <td style="width:20%;vertical-align:top;">{cell("❤️","Customer Relationships","customer_relationships")}</td>
  <td rowspan="2" style="width:20%;vertical-align:top;">{cell("👥","Customer Segments","customer_segments")}</td>
</tr>
<tr>
  <td style="vertical-align:top;">{cell("🏗️","Key Resources","key_resources")}</td>
  <td style="vertical-align:top;">{cell("🚚","Channels","channels")}</td>
</tr>
<tr>
  <td colspan="2" style="vertical-align:top;">{cell("🏷️","Cost Structure","cost_structure", bot_style)}</td>
  <td></td>
  <td colspan="2" style="vertical-align:top;">{cell("💵","Revenue Streams","revenue_streams", bot_style)}</td>
</tr>
</table>
""", unsafe_allow_html=True)

    ue = result.get("unit_economics", {})
    if ue:
        st.markdown("### Unit economics (rough estimate)")
        ue_rows = [("Price", "price"), ("Unit cost", "unit_cost"), ("Gross margin", "gross_margin"),
                   ("Customer acquisition cost", "customer_acquisition_cost"), ("Payback", "payback")]
        st.dataframe(pd.DataFrame([{"Metric": n, "Estimate": ue.get(k, "—")} for n, k in ue_rows]),
                     hide_index=True, use_container_width=True)
        if ue.get("note"):
            st.caption(f"Biggest unknown: {ue.get('note')}")

# ---------------- tab 7: MRD / PRD
with tab7:
    st.markdown("### Write an MRD or PRD")
    st.write("Turn this analysis into a first draft you can edit. "
             "Anything that depends on an untested assumption is marked (to validate).")
    dtype = st.radio("Document", ["MRD", "PRD"], horizontal=True,
                     captions=["Market Requirements: who the market is and what it needs",
                               "Product Requirements: what to build first and how to measure it"])
    if dtype == "PRD" and str(result.get("confidence", "")).lower().startswith("low"):
        st.info("Confidence is still low, so many requirements will be marked (to validate). "
                "Many teams write the MRD first and the PRD after the first tests.")
    test_results = st.text_area("Test results so far (optional)", height=100,
                                placeholder="e.g. 7 of 9 interviewees track product expiry manually; "
                                            "survey: acceptable price range $199-$299")
    tpl_file = st.file_uploader("Use your own template (optional): DOCX, MD, TXT, or PDF",
                                type=["docx", "md", "txt", "pdf"], key="tpl_upload")
    template = MRD_TEMPLATE if dtype == "MRD" else PRD_TEMPLATE
    if tpl_file:
        try:
            template = read_template(tpl_file)
            st.success("Your template will be used.")
        except Exception as e:
            st.error(f"Could not read the template: {e}")
    with st.expander("See the template structure"):
        st.markdown(template)

    if st.button(f"Write {dtype} draft", type="primary"):
        with st.spinner(f"Writing your {dtype}... about 40-60 seconds."):
            try:
                st.session_state["doc"] = {"type": dtype, "text": generate_doc(
                    result, st.session_state["history"][-1]["idea"], dtype, template, test_results)}
            except Exception as e:
                st.error(f"Request failed: {e}")

    doc = st.session_state.get("doc")
    if doc:
        with st.container(border=True):
            st.markdown(doc["text"])
        c1, c2, _ = st.columns([1.3, 1.3, 5])
        try:
            c1.download_button(f"Download {doc['type']} (Word)", markdown_to_docx(doc["text"]),
                               file_name=f"{doc['type']}_draft.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        except Exception:
            c1.caption("Word export failed for this draft. Use Markdown instead.")
        c2.download_button(f"Download {doc['type']} (Markdown)", doc["text"],
                           file_name=f"{doc['type']}_draft.md", mime="text/markdown")